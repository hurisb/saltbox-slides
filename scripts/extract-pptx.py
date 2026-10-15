#!/usr/bin/env python3
"""
Extract text, speaker notes and images from a PowerPoint so Claude can rebuild it as a Saltbox deck.

    pip install python-pptx
    python3 scripts/extract-pptx.py deck.pptx extracted/

Writes:
  extracted/outline.md   one section per slide: title, text blocks, notes, images found
  extracted/media/       every picture, named slide-03-image-1.png, …
"""
import pathlib, sys

try:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
except ImportError:
    sys.exit("python-pptx is missing. Run:  pip install python-pptx")

if len(sys.argv) < 2:
    sys.exit(__doc__)
src = pathlib.Path(sys.argv[1])
out = pathlib.Path(sys.argv[2] if len(sys.argv) > 2 else "extracted")
media = out / "media"
media.mkdir(parents=True, exist_ok=True)

prs = Presentation(src)
lines = [f"# {src.name}", f"{len(prs.slides)} slides\n"]

def walk(shapes):
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from walk(sh.shapes)
        else:
            yield sh

for n, slide in enumerate(prs.slides, 1):
    title = slide.shapes.title.text.strip() if slide.shapes.title is not None and slide.shapes.title.has_text_frame else ""
    lines.append(f"## Slide {n}: {title or '(no title)'}")
    img = 0
    for sh in walk(slide.shapes):
        if sh.has_text_frame and sh != slide.shapes.title:
            text = "\n".join(p.text.strip() for p in sh.text_frame.paragraphs if p.text.strip())
            if text:
                lines.append("- " + text.replace("\n", "\n  - "))
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
            img += 1
            name = f"slide-{n:02d}-image-{img}.{sh.image.ext}"
            (media / name).write_bytes(sh.image.blob)
            lines.append(f"- [image] media/{name} ({sh.width // 9525}×{sh.height // 9525}px on slide)")
        if getattr(sh, "has_table", False) and sh.has_table:
            for row in sh.table.rows:
                lines.append("- | " + " | ".join(c.text.strip() for c in row.cells) + " |")
    if slide.has_notes_slide:
        notes = slide.notes_slide.notes_text_frame.text.strip()
        if notes:
            lines.append("\n**Notes:** " + notes.replace("\n", " "))
    lines.append("")

(out / "outline.md").write_text("\n".join(lines))
print(f"✓ {len(prs.slides)} slides → {out / 'outline.md'}  ·  images → {media}")
